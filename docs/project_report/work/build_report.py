from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import json, hashlib, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
FIG=ROOT/'docs/research_paper/work/figures'
OUT=HERE.parent/'OpthaEdge_Engineering_Project_Report.docx'
FONT=Path('C:/Windows/Fonts')
def ft(s,b=False):return ImageFont.truetype(str(FONT/('arialbd.ttf' if b else 'arial.ttf')),s)
def node(dr,xy,title,lines=(),fill='#eaf1f6'):
    x,y,w,h=xy;dr.rounded_rectangle((x,y,x+w,y+h),radius=15,fill=fill,outline='#36556c',width=4)
    dr.text((x+w/2,y+22),title,font=ft(36,True),fill='#152c3b',anchor='ma')
    for i,t in enumerate(lines):dr.text((x+w/2,y+75+i*42),t,font=ft(31),fill='#243b4b',anchor='ma')
def arrow(dr,a,b):
    dr.line([a,b],fill='#36556c',width=4);ang=math.atan2(b[1]-a[1],b[0]-a[0]);sz=18
    dr.polygon([b,(b[0]-sz*math.cos(ang-.45),b[1]-sz*math.sin(ang-.45)),(b[0]-sz*math.cos(ang+.45),b[1]-sz*math.sin(ang+.45))],fill='#36556c')

im=Image.new('RGB',(2100,1230),'white');dr=ImageDraw.Draw(im)
xs=[200,740,1290,1860]
for x,t in zip(xs,['Browser','Prediction route','Inference module','ONNX session']):
    node(dr,(x-180,20,360,110),t)
    dr.line((x,130,x,1120),fill='#b4bdc4',width=3)
messages=[(0,1,230,'1  multipart image + mode'),(1,1,340,'2  file and decode checks'),(1,2,465,'3  call selected task'),(2,3,585,'4  acquire cached session'),(2,2,705,'5  resize and normalize'),(2,3,815,'6  tensor execution'),(3,2,910,'7  output logits'),(2,1,1000,'8  validate and decode'),(1,0,1090,'9  result + model metadata')]
for a,b,y,label in messages:
    if a==b:
        dr.line([(xs[a],y),(xs[a]+140,y),(xs[a]+140,y+48)],fill='#36556c',width=4);arrow(dr,(xs[a]+140,y+48),(xs[a],y+48));dr.text((xs[a]+160,y+5),label,font=ft(29),fill='#152c3b')
    else:
        arrow(dr,(xs[a],y),(xs[b],y));dr.text(((xs[a]+xs[b])/2,y-38),label,font=ft(29),fill='#152c3b',anchor='ma')
dr.text((1050,1190),'Validation or execution failure returns a structured error instead of a screening result.',font=ft(30),fill='#243b4b',anchor='ma')
im.save(HERE/'sequence.png',dpi=(300,300))

im=Image.new('RGB',(1800,1400),'white');dr=ImageDraw.Draw(im)
node(dr,(420,20,960,170),'Browser and authorized user',['Research or approved clinical workflow'])
node(dr,(420,270,960,205),'Web application and identity boundary',['OIDC, session validation and resource authorization','Model approval and request limits'])
arrow(dr,(900,190),(900,270))
node(dr,(70,610,760,210),'Patient data services',['Organization-scoped records and audit','PostgreSQL transactions'])
node(dr,(970,610,760,210),'Isolated inference worker',['Queue, timeouts and resource limits','Quality gate and approved model versions'])
arrow(dr,(650,475),(450,610));arrow(dr,(1150,475),(1350,610))
node(dr,(70,980,760,200),'Private image storage',['Encrypted originals and derived assets','Content hashes and controlled access'])
node(dr,(970,980,760,200),'Operational controls',['Monitoring, backups and recovery tests','Versioned releases and rollback'])
arrow(dr,(1350,820),(450,980));arrow(dr,(450,820),(450,980));arrow(dr,(1350,820),(1350,980))
dr.text((900,1300),'PROPOSED TARGET: these integrations are not established by the current prototype.',font=ft(34,True),fill='#243b4b',anchor='ma')
im.save(HERE/'deployment.png',dpi=(300,300))

d=Document()
for sec in d.sections:
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.left_margin=Inches(1.25);sec.right_margin=Inches(1);sec.top_margin=Inches(1);sec.bottom_margin=Inches(1)
    sec.header_distance=Inches(.45);sec.footer_distance=Inches(.45)
styles=d.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC Heading']:
    s=styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if False else None
normal=styles['Normal'];normal.font.size=Pt(12);normal.font.bold=False;normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',24),('Subtitle',16),('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
    styles[name].font.size=Pt(size);styles[name].font.bold=name!='Subtitle';styles[name].paragraph_format.keep_with_next=True
    styles[name].paragraph_format.space_before=Pt(10);styles[name].paragraph_format.space_after=Pt(9)
styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False;styles['Caption'].paragraph_format.space_after=Pt(8)
for name in ['TOC 1','TOC 2']:
    if name not in styles:styles.add_style(name,1)
    styles[name].font.name='Times New Roman';styles[name].font.size=Pt(11);styles[name].paragraph_format.space_after=Pt(4)

for s in styles:
    for border in s.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    if s.type in (1,2):
        rf=s.element.get_or_add_rPr().find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');s.element.get_or_add_rPr().insert(0,rf)
        rf.attrib.clear()
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
styles['Subtitle'].font.italic=False

figures=[];tables=[]
def p(t='',boldlead=None):
    e=d.add_paragraph();e.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if boldlead and t.startswith(boldlead):e.add_run(boldlead).bold=True;e.add_run(t[len(boldlead):])
    else:e.add_run(t)
    return e
def head(t,level=1,new=True):
    e=d.add_paragraph(t,f'Heading {level}')
    if new:e.paragraph_format.page_break_before=True
    return e
def minor(t):return head(t,2,False)
def bullet(t):
    e=d.add_paragraph(t,'List Bullet');e.paragraph_format.space_after=Pt(5);return e
def code(t):
    e=d.add_paragraph();e.paragraph_format.space_after=Pt(5);e.paragraph_format.line_spacing=1
    r=e.add_run(t);r.font.name='Consolas';r.font.size=Pt(9)
    return e
def field(e,instruction):
    r=e.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    r=e.add_run();a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=instruction;r._r.append(a)
    r=e.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)

def bookmark(e,name,ident):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(ident));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(ident))
    e._p.insert(0,start);e._p.append(end)
def image(path,title,width=5.6):
    e=d.add_paragraph();e.alignment=WD_ALIGN_PARAGRAPH.CENTER;e.paragraph_format.keep_with_next=True;e.paragraph_format.space_after=Pt(3)
    e.add_run().add_picture(str(path),width=Inches(width));n=len(figures)+1
    e=d.add_paragraph(f'Figure {n}  {title}','Caption');e.paragraph_format.keep_together=True;bookmark(e,f'fig{n}',100+n);figures.append((n,title));return n
def table(title,headers,rows,widths=None):
    n=len(tables)+1;e=d.add_paragraph(f'Table {n}  {title}','Caption');e.paragraph_format.keep_with_next=True;bookmark(e,f'tab{n}',200+n);tables.append((n,title))
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [6.0/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for side,v in [('top','65'),('bottom','65'),('left','80'),('right','80')]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    for ri,vals in enumerate([headers]+rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,value) in enumerate(zip(cells,vals)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=1
            pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.05
            rr=pp.add_run(str(value));rr.font.size=Pt(10.5);rr.bold=ri==0
            if ri==0:rr.font.color.rgb=RGBColor(255,255,255)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'36556C' if ri==0 else ('F3F6F8' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(sh)
    spacer=p('');spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3);spacer.runs[0].font.size=Pt(1)
    return t
def footer(sec,fmt,start):
    sec.footer.is_linked_to_previous=False;sec.header.is_linked_to_previous=False
    pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),fmt);pg.set(qn('w:start'),str(start));sec._sectPr.append(pg)
    e=sec.footer.paragraphs[0];e.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(e,' PAGE ')
    for r in e.runs:r.font.size=Pt(10)
    e=sec.header.paragraphs[0];e.text='OPTHAEDGE  |  ENGINEERING PROJECT REPORT';e.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in e.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('525D66')
def eq(parts):
    e=d.add_paragraph();e.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath')
    def mr(t):r=OxmlElement('m:r');v=OxmlElement('m:t');v.text=t;r.append(v);return r
    def append(parent,part):
        if isinstance(part,list):
            for item in part:append(parent,item)
        elif isinstance(part,tuple):
            tag,a_tag,b_tag=('m:sSub','m:e','m:sub') if part[0]=='sub' else ('m:f','m:num','m:den')
            q=OxmlElement(tag);a=OxmlElement(a_tag);b=OxmlElement(b_tag)
            append(a,part[1]);append(b,part[2]);q.extend([a,b]);parent.append(q)
        else:parent.append(mr(part))
    append(om,parts)
    e._p.append(om)

# Cover
e=d.add_paragraph('ENGINEERING PROJECT REPORT');e.alignment=1;e.paragraph_format.space_before=Pt(55);e.runs[0].font.size=Pt(14)
e=d.add_paragraph('OpthaEdge','Title');e.alignment=1;e.paragraph_format.space_before=Pt(28)
e=d.add_paragraph('Diabetic Retinopathy Research Screening\nand Model Governance','Subtitle');e.alignment=1
p('').paragraph_format.space_after=Pt(35)
for t in ['Prepared by','Shamail Pawar','Faizan Mustafa Shaikh','Ayaan Shaikh']:
    e=d.add_paragraph(t);e.alignment=1;e.paragraph_format.space_after=Pt(5)
e=d.add_paragraph('Project guide\nNayana Mahajan');e.alignment=1;e.paragraph_format.space_before=Pt(24)
e=d.add_paragraph('Department of Electronics & Computer Science Engineering\nVidyalankar Institute of Technology\nMumbai, India');e.alignment=1;e.paragraph_format.space_before=Pt(28)
e=d.add_paragraph('Source and verification baseline\n29 September 2026');e.alignment=1;e.paragraph_format.space_before=Pt(30)

sec=d.add_section(WD_SECTION_START.NEW_PAGE);footer(sec,'lowerRoman',1)
head('Abstract',1,False)
p('OpthaEdge is a research-stage web application for reviewing diabetic retinopathy model outputs from retinal fundus photographs. The project combines two local inference pipelines, a model-information registry, a browser-based screening interface and a separate foundation for organization-scoped patient records. This report documents the implemented engineering system, its verification evidence and the work required before a clinical deployment could be considered.')
p('The application is built with Next.js, React and TypeScript. A Node.js route validates uploaded JPEG or PNG content, invokes task-specific Sharp preprocessing and runs an ONNX model on the CPU. The binary path uses a 224 × 224 input and two logits, while the severity path uses a 299 × 299 input and five logits. Output checks reject unexpected element counts and non-finite values before stable softmax normalization. Responses include model and preprocessing versions. Patient services use development sessions, permission checks, organization boundaries and transactional audit writes through Prisma and PostgreSQL.')
p('The recorded verification baseline contains 12 passing unit tests and nine passing local HTTP checks, together with successful type checking and linting. Two HTTP checks executed the models on a generated constant-color image. These checks demonstrate selected software behavior; they do not measure diagnostic accuracy, clinical usefulness or deployment throughput. Training provenance, verified class order, validated thresholds, calibration and independent clinical evaluation remain unavailable.')
p('The engineering outcome is an inspectable research integration with explicit implementation boundaries. This report provides requirements, architecture and sequence diagrams, a data model, preprocessing equations, API contracts, project screenshots, test traceability, operating instructions and a staged development plan.')
p('Keywords: diabetic retinopathy, ONNX, image preprocessing, web application, model governance, software verification.')
minor('Acknowledgment')
p('The authors acknowledge their project guide, Nayana Mahajan, for guidance on this work. The report identifies external research and documentation in the reference section and distinguishes that literature from evidence produced by the project.')

e=head('Contents',1,True);e.style='TOC Heading'
toc=d.add_paragraph();field(toc,' TOC \\o "1-1" \\h \\z \\u ')

head('List of figures and tables')
minor('List of figures');fig_list=d.add_paragraph('__FIGURES__')
minor('List of tables');table_list=d.add_paragraph('__TABLES__')
# Abbreviations separate page avoids tiny front-matter text.
head('Abbreviations and report conventions')
table('Abbreviations',['Term','Meaning'],[['API','Application programming interface'],['AUROC','Area under the receiver operating characteristic curve'],['CPU','Central processing unit'],['DR','Diabetic retinopathy'],['DTO','Data transfer object'],['MIME','Multipurpose Internet Mail Extensions'],['OIDC','OpenID Connect'],['ONNX','Open Neural Network Exchange'],['ORM','Object-relational mapper'],['QWK','Quadratic weighted kappa'],['RBAC','Role-based access control'],['RGB','Red, green and blue image channels'],['SHA-256','Cryptographic digest used for artifact identity']],[1.0,5.0])
p('Implemented means represented in the reviewed application source. Recorded verification means supported by the retained test evidence dated 29 September 2026. Schema only means declared in Prisma without a complete active screening workflow. Proposed means future engineering work, not an existing capability.')
p('The word probability in an API field describes a normalized model score. It does not establish clinical calibration. Clinical accuracy and production suitability are not inferred from successful software tests. Unless stated otherwise, the report concerns source commit 0099145dc14547c30484798042b029fbf58df1eb.')

sec=d.add_section(WD_SECTION_START.NEW_PAGE);footer(sec,'decimal',1)
head('1 Introduction',1,False)
minor('1.1 Background and motivation')
p('Diabetic retinopathy affects retinal blood vessels and can cause vision loss. Early disease may be asymptomatic, which makes systematic eye assessment important [1]. Research on automated analysis of retinal photographs motivates tools that can support image review, but the availability of an executable model is only one part of a dependable screening workflow.')
p('An engineering system must preserve the meaning of its inputs and outputs. It needs clear image-validation rules, reproducible preprocessing, identified model artifacts, predictable error handling and an interface that communicates the limits of model evidence. If patient information is involved, identity, consent, authorization, storage and auditing also become essential concerns.')
minor('1.2 Problem statement')
p('The project integrates two exported diabetic-retinopathy models whose training and clinical-validation packages have not been supplied. The engineering problem is to expose their inference contracts through a usable research application while preventing implementation details, database plans or model scores from being mistaken for evidence of clinical validity. A second problem is to establish patient-record services without implying that the screening route already creates a longitudinal clinical record.')
minor('1.3 Objectives')
for t in ['Implement separate binary and five-level severity inference paths with explicit preprocessing and output validation.','Provide a browser workflow for task selection, upload, result inspection and structured failure messages.','Associate each response with artifact and preprocessing identities, and make missing model evidence visible.','Provide a development patient-record foundation with permission checks, organization scoping and audit writes.','Verify selected software boundaries and document the remaining engineering and evaluation work.']:bullet(t)
minor('1.4 Scope')
p('The report covers the existing web interface, local CPU inference, model registry, patient data-access foundation and declared database schema. It does not describe model training, clinical trial results, validated lesion localization, macular-edema detection or autonomous referral. The name OpthaEdge does not establish on-device or low-power edge performance; inference currently runs in a Node.js server process.')

head('2 Literature and technical background')
p('Published diabetic-retinopathy studies provide context for the application, but their results are not transferable to the supplied artifacts. The distinction is important because model architecture, preprocessing and label names do not identify the dataset, reference standard or evaluation protocol used to develop a particular model file.')
table('Relevant literature and engineering implications',['Source','Contribution','Implication for OpthaEdge'],[['Gulshan et al. [2]','Deep-learning evaluation for referable DR with retinal photographs.','Independent reference-standard data are required.'],['Ting et al. [3]','Evaluation across multiethnic retinal-image datasets.','Assess population and acquisition-device variation.'],['Abramoff et al. [4]','Prospective assessment of an autonomous DR system in primary care.','A deployable workflow needs prospective evidence.'],['Szegedy et al. [5]','Inception architecture design.','Architecture identity does not prove artifact provenance.'],['Guo et al. [6]','Calibration of neural-network scores.','Softmax scores require calibration assessment.'],['CLAIM and TRIPOD+AI [7, 8]','Transparent reporting of imaging AI and prediction-model studies.','Prespecify data, methods, metrics and limitations.']],[1.25,2.30,2.45])
minor('2.1 Technical foundations')
p('A convolutional image classifier transforms image tensors into output scores. In this application, the exported artifacts are consumed through ONNX Runtime, which supplies the Node.js inference binding [9]. The application does not contain a reproducible training pipeline, dataset manifest or learned-weight development history. Consequently, the report describes the observed integration contract rather than reconstructing an unsupported internal network architecture.')
p('Stable softmax converts logits into nonnegative values that sum to one before display rounding. It preserves relative score ordering but does not prove that a reported value reflects the likelihood of clinical correctness. Threshold selection, label ordering and calibration must be evaluated for the exact artifact and intended population [6].')
minor('2.2 Project contribution')
p('OpthaEdge contributes a source-inspectable integration: explicit task branching, file validation, versioned model metadata, structured errors and a visible separation between working research features and deferred clinical services. No superiority claim or numerical comparison with the cited diagnostic systems is made.')

head('3 Requirements and feasibility')
minor('3.1 Actors and operating assumptions')
p('The research operator selects a task and submits an image for technical review. A development user with patient permissions can search or create synthetic patient records. A maintainer configures the runtime and inspects model metadata. Production clinical roles are design targets; synthetic development sign-in does not establish a production identity system.')
table('Functional requirements and implementation status',['ID','Requirement','Current status'],[['FR-01','Select binary or severity task explicitly.','Implemented in client and API.'],['FR-02','Reject invalid file content, type, size and unreadable dimensions.','Implemented; retinal quality is not assessed.'],['FR-03','Apply the preprocessing contract for the selected task.','Implemented; training correspondence unverified.'],['FR-04','Execute the local artifact and validate output values.','Implemented for float32, count and finite values.'],['FR-05','Return model metadata and structured success or error.','Implemented discriminated response union.'],['FR-06','Expose model limitations and deployment status.','Implemented registry and information page.'],['FR-07','List and create organization-scoped patient records.','Implemented development DAL and server action.'],['FR-08','Persist screening images and results to an examination.','Schema foundation only.'],['FR-09','Approve and sign clinical reviews with amendments.','Deferred; route disabled.'],['FR-10','Enforce authorization, consent and model approval on inference.','Required future integration.']],[.62,3.10,2.28])
minor('3.2 Input and output requirements')
p('The prediction request contains a File entry named image and a mode value equal to binary or severity. The accepted byte formats are JPEG and PNG. The route uses a 10 MiB upload limit and a 50-million-pixel decode limit. A successful response contains a task-specific result, elapsed inference time and model-run metadata. Failed requests return a machine-readable error code and a human-readable message.')
p('Patient identifiers use 3-32 uppercase letters, numbers or hyphens. The server supplies organization scope; client forms cannot select another tenant.')

head('3.3 Nonfunctional requirements',2,True)
table('Nonfunctional requirements and remaining evidence',['ID','Requirement','Evidence or gap'],[['NFR-01','Reproducible model execution','Registry identifies artifacts and preprocessing; runtime digest enforcement is absent.'],['NFR-02','Fail predictably on malformed input','Selected unit and HTTP checks passed; transport and adversarial testing remain.'],['NFR-03','Protect patient boundaries','DAL checks permission and organization; complete database isolation testing is not recorded.'],['NFR-04','Maintain understandable result semantics','Research warnings are displayed; scores remain uncalibrated.'],['NFR-05','Control resource consumption','File/pixel limits and cached sessions exist; queue, timeout and load policies are missing.'],['NFR-06','Operate and recover reliably','Versioned source exists; monitoring, recovery and rollback exercises are not recorded.']],[.70,2.05,3.25])
minor('3.4 Feasibility assessment')
p('Technical feasibility is supported at prototype level by successful artifact execution and the existing browser workflow. The Node.js deployment model accommodates ONNX Runtime, Sharp and local artifact access. This does not establish feasibility on a constrained edge device, a serverless execution limit or an untested operating system.')
p('Operational feasibility remains conditional. The patient-record foundation provides a useful separation of concerns, but clinical operation would require approved identity, storage, quality assessment, model governance and review processes. A supported screening workflow must define responsibility for unreadable images, missing results and disagreement between an operator and model output.')
p('Economic feasibility has not been measured. No cloud bill, maintenance budget, annotation cost or clinician-workload study is available. Local execution may avoid a separate inference API charge, but model maintenance, compute, secure storage, backups, external evaluation and support still create costs. The report therefore gives dependencies rather than invented expenditure estimates.')
minor('3.5 Development hardware and software')
p('The retained HTTP verification record identifies Windows, Node.js v25.5.0 and an Intel Core i7-10850H processor with 12 logical processors. These describe the recorded environment, not minimum specifications. The README requires PostgreSQL 15 or newer and a Node.js release supported by the pinned Prisma dependencies. Reproducibility should begin with the lockfile and an approved supported Node release; RAM, GPU, throughput and cold-start requirements have not been benchmarked.')

head('4 System design')
minor('4.1 Implemented architecture')
p('Figure 1 separates research inference from patient-record services. The client screening component calls a Node.js route that uses local model files. Patient requests instead pass through session-derived authorization and a server-only data-access layer. The screening route does not currently enter the patient transaction or persist an examination.')
image(FIG/'architecture.png','Implemented architecture and planned services',4.4)
p('The model loader caches both a successfully created session and an in-progress initialization promise. Concurrent first-use callers can share initialization within one process. Failed initialization clears the promise so a later call can retry. This mechanism does not provide a global cache, workload queue or resource limit across multiple server processes.')

head('4.2 Request sequence and data flow',2,True)
p('Figure 2 follows the successful inference path. Validation failures stop processing before model execution where possible. Execution or decoding failures produce structured errors. Input file validation concerns technical content and size; it does not classify retinal quality or clinical suitability.')
image(HERE/'sequence.png','Sequence of a research prediction request',6.0)
minor('Data-flow boundaries')
table('Data flows and persistence boundaries',['Flow','Data','Current handling'],[['Browser to route','Multipart file and selected mode','Received and parsed by Node.js route.'],['Route to preprocessing','Image bytes','Decoded and resized in server memory.'],['Preprocessing to model','Float32 channel-first tensor','Passed to the selected CPU session.'],['Model to decoder','Two or five logits','Count and finite-value checks; stable softmax.'],['Route to browser','Result, elapsed time, model metadata','Stored in client component state.'],['Patient action to DAL','Validated patient fields plus server actor','Transactional database writes and audit.']],[1.55,1.75,2.70])
p('No screening-to-patient association is present in this path. Browser refresh discards the transient screening state. A future persistent workflow must introduce explicit examination identity, consent linkage and immutable model-version association; it cannot infer these from an uploaded filename.')

head('4.3 Database design',2,True)
p('The Prisma schema models identity, patient records, examinations, image metadata, model versions and predictions. Figure 3 shows a selected relationship path; it is not an exhaustive entity diagram. Blue elements represent development patient-record functionality, while gray elements express screening structures that are not populated by the prediction route.')
image(FIG/'data_model.png','Selected data relationships and implementation boundaries',4.5)
p('Patient numbers are unique within an organization. Image records include laterality, a storage key and a content digest. A prediction links an image to a model version and associated per-label scores. Binary image content is intended for private object storage rather than direct database storage, but that storage integration is not established.')

head('4.4 Data dictionary and authorization',2,True)
table('Selected entities and integrity responsibilities',['Entity','Representative data','Responsibility'],[['Organization','Identity, status and membership scope','Tenant boundary.'],['User and membership','OIDC subject, active period, role','Construct current authorization actor.'],['Patient','Organization, patient number, status','Stable organization-scoped record.'],['Patient demographics','Name, birth date, sex, language','Separate personal attributes.'],['ConsentRecord','Purpose, status, policy version, actor','Document permission context.'],['Examination','Patient, organization, encounter context','Future screening association.'],['EyeImage','Examination, eye, storage key, SHA-256','Reference an original captured image.'],['ModelVersion','Artifact hash, labels, approval metadata','Identify the model behind a prediction.'],['Prediction and probability','Image/model links and label scores','Future immutable inference history.'],['AuditEvent','Actor, action, resource, outcome, time','Record selected patient operations.']],[1.45,2.40,2.15])
minor('Session and patient access')
p('Development sessions are signed with HS256, validate issuer and audience, and expire after eight hours. The cookie is HTTP-only and SameSite=Lax; secure transport is required by the production cookie setting. Synthetic authentication is enabled only outside production and only when explicitly configured. A session secret shorter than 32 characters is rejected.')
p('The server rebuilds the current actor from an active user and active organization memberships. Patient listing requires read permission, filters by that actor’s organization, excludes archived records and returns at most 50 minimal DTOs. Patient creation writes the patient, selected consents and successful audit event in one database transaction. A pre-insert duplicate lookup improves feedback, while the schema’s unique key remains necessary under concurrent requests.')
p('These controls do not cover the public prediction endpoint. They also do not prove complete isolation, secure backup, retention enforcement or append-only audit behavior. Those properties require integration tests and operational mechanisms beyond the current helper-level test suite.')

head('5 Inference methodology')
minor('5.1 Separate task pipelines')
p('The system treats binary screening and severity grading as separate contracts. Figure 4 shows their preprocessing and decoding branches. This avoids silently applying one transformation to artifacts that expect different spatial dimensions and normalization. Both transformations remain implementation assumptions until reconciled with training and export records.')
image(FIG/'inference.png','Task-specific preprocessing and output decoding',4.5)
p('Sharp uses cover resizing and removes alpha before reading raw pixel values. Cover resizing can crop image edges. The functions assume three color channels but do not explicitly convert every accepted input color space to RGB. Consequently, accepting a JPEG or PNG does not guarantee a valid three-channel tensor for all grayscale or unusual-color-space images.')

head('5.2 Mathematical model and algorithms',2,True)
p('Let v be an 8-bit pixel value and c a color-channel index. Binary preprocessing scales and normalizes each channel as follows, where μ = (0.485, 0.456, 0.406) and σ = (0.229, 0.224, 0.225).')
eq([('sub','x','c'),' = ',('frac',['v / 255 − ',('sub','μ','c')],('sub','σ','c'))])
p('The severity path uses x = v / 127.5 − 1. Both paths then arrange data in channel-first order. A single binary tensor contains 150,528 float32 values, or 602,112 raw bytes; the severity tensor contains 268,203 values, or 1,072,812 raw bytes. These calculated sizes exclude model weights, decoder buffers, runtime workspaces and concurrent requests.')
p('For logits z and maximum logit m, the decoder computes stable softmax. Subtracting m reduces numerical overflow without changing the normalized relative scores.')
eq([('sub','p','k'),' = ',('frac',['exp(',('sub','z','k'),' − m)'],[('sub','Σ','j'),' exp(',('sub','z','j'),' − m)']),',    m = ',('sub','max','j'),' ',('sub','z','j')])
p('Binary decoding returns DR Detected when p₁ ≥ 0.5. The API probability field is the DR score p₁ even when No DR is selected. The confidence field is 100 times the selected class score. Neither is a calibrated clinical confidence estimate. Severity uses the index of the largest score and returns a weighted index s, which is a score summary rather than a validated clinical measurement.')
eq(['s = 0p₀ + 1p₁ + 2p₂ + 3p₃ + 4p₄'])
minor('Algorithm 1 Research inference')
for t in ['Parse multipart data; require a File and an explicit supported mode.','Reject empty or oversized content; identify JPEG/PNG signature and validate declared MIME.','Read dimensions with the configured pixel limit; stop if content cannot be decoded.','Acquire the selected cached CPU session and construct the task-specific tensor.','Execute the first declared model input/output pair; require float32 output.','Validate element count and finite values; normalize and decode in registry order.','Return the result, elapsed time and model metadata, or a structured error.']:bullet(t)
p('Display rounding occurs after decoding. Returned severity scores may not sum to exactly one after rounding. Element-count validation is not complete tensor-axis validation; unexpected rank and channel handling remain future test targets.')

head('6 Implementation and API design')
table('Implementation module map',['Module','Main responsibility'],[['components/ScreeningWorkspace.tsx','Task selection, upload, request state and result review.'],['app/api/predict/route.ts','Multipart contract, file validation, dispatch and error mapping.'],['lib/imageFileSignature.ts','JPEG/PNG signature and MIME consistency helpers.'],['lib/preprocessor.ts','Task-specific resizing, normalization and channel ordering.'],['lib/modelLoader.ts','Local artifacts and cached ONNX CPU sessions.'],['lib/inference.ts','Session execution and result assembly.'],['lib/outputDecoding.ts','Logit validation, stable softmax and output semantics.'],['lib/modelRegistry.ts','Model identities, versions, contracts and limitations.'],['lib/auth and data/patients.ts','Development sessions, authorization and patient DAL.'],['prisma/schema.prisma','Normalized identity, patient and screening schema.']],[2.50,3.50])
minor('6.1 Technology choices')
p('Next.js 16.2.1 supplies the application and Node.js request handler; React 19.2.4 supports the interface. TypeScript expresses task-specific response unions. Sharp performs image decoding and tensor preparation, while ONNX Runtime provides local execution [9]. Prisma and PostgreSQL support relational records; Zod validates patient input; jose handles signed development sessions. Versions with range specifiers must be resolved through package-lock.json for an exact installation.')
p('The prediction route explicitly selects the Node.js runtime. Native inference libraries and local artifact reads require a compatible server environment; a browser-only or Edge-runtime deployment is not demonstrated. Production build scripts regenerate Prisma Client before compiling the application.')
minor('6.2 API request and result contract')
code('POST /api/predict\nContent-Type: multipart/form-data\nimage: one JPEG or PNG File\nmode: binary | severity')
p('A successful response has success=true and a result containing mode, the relevant binary or severity payload, inferenceTimeMs and model metadata. An error has success=false and an error object containing code and message. Consumers must branch on success and mode instead of assuming every response contains every result field.')

head('6.3 Error handling and model governance',2,True)
table('Prediction API error taxonomy',['Condition','Error code','HTTP'],[['Malformed multipart body','INVALID_FORM_DATA','400'],['File missing or wrong field type','MISSING_IMAGE','400'],['Unsupported task','INVALID_MODE','400'],['Zero-byte upload','EMPTY_FILE','400'],['Upload over 10 MiB','FILE_TOO_LARGE','413'],['Unsupported content signature','UNSUPPORTED_IMAGE_CONTENT','415'],['Declared type differs from content','MIME_CONTENT_MISMATCH','400'],['Decode or metadata failure','CORRUPT_IMAGE','400'],['Explicit dimension-limit branch','IMAGE_DIMENSIONS_TOO_LARGE','413'],['Invalid logit count or values','MODEL_OUTPUT_INVALID','503'],['Other execution failure','INFERENCE_FAILED','500']],[2.20,3.13,.67])
p('The error taxonomy describes source branches, not proof that each branch was exercised. Sharp may reject excessive pixel dimensions while reading metadata, in which case the catch block returns CORRUPT_IMAGE. A missing or non-float32 output tensor raises a generic execution error, whereas decoder count/finite failures use MODEL_OUTPUT_INVALID. These distinctions matter when specifying additional regression tests.')
minor('6.4 Registry semantics')
p('Each registry entry identifies the task, artifact file, SHA-256 version, input dimensions, assumed output labels and preprocessing version. Both entries explicitly record unverified label order, threshold and training correspondence, with no validation dataset or metrics. The clinical deployment status is blocked-pending-validation.')
p('This status is descriptive metadata. The API currently does not enforce approval, session authentication, consent or patient authorization before inference. Registry hashes matched the artifacts in the recorded assessment, but the model loader does not recompute and enforce those digests. A future release must turn these recorded assumptions into enforced checks where appropriate.')

head('7 User interface and workflow')
p('The following images are authentic captures from the local application used for the research paper. They document implemented screens and contain no patient photographs, identifiable patient information or clinical prediction examples. Figure 5 presents task selection; Figure 6 shows the upload interface.')
image(FIG/'screening.png','Research task selection in the actual application',6.0)
image(FIG/'upload.png','Actual image-upload interface',6.0)

head('7.1 Model information and user procedure',2,True)
image(FIG/'models.png','Model-information page and validation blockers',6.0)
p('Figure 7 exposes the binary model entry and validation limitations. The lower severity entry is outside this captured viewport. The status badge communicates missing evidence; it does not by itself disable the research endpoint.')
minor('Research screening procedure')
for t in ['Open the screening workspace and choose binary DR or DR severity grading.','Select a permitted research image and review the upload state before submission.','Start analysis and wait for the request to complete; inspect either the task result or the structured error.','Read the model identity and research limitations with the output. Do not interpret a score as a clinical referral decision.','Change the selected task to clear the existing file and result; refresh also discards transient screening state.']:bullet(t)
minor('Patient workflow boundary')
p('Synthetic development sign-in and patient creation are separate from screening. The patient list depends on a configured database, generated client, migrations, seed data and the development authentication settings. Examinations and reports show requirements rather than a complete screening workflow. Review Queue and Follow-ups route files exist but return not-found while their module flags are false.')

head('8 Testing and verification')
minor('8.1 Verification strategy and provenance')
p('The verification baseline was recorded on 29 September 2026 for commit 0099145dc145. The current report reuses the retained evidence and reviews the corresponding source; it does not claim a new test run. Verification combines helper-level unit tests, local HTTP checks, type checking and linting. No patient cohort, reference-standard grading or clinical outcome data were used.')
table('Existing unit-test groups',['Test file','Cases','What is exercised'],[['authPolicy.test.ts','2','Permission membership and cross-organization rejection.'],['imageFileSignature.test.ts','4','JPEG/PNG signatures, invalid/truncated signatures and MIME matching.'],['outputDecoding.test.ts','4','Two/five-logit decoding, unexpected counts and non-finite values.'],['patientValidation.test.ts','2','Patient-ID normalization and invalid identifier/name rejection.']],[2.10,.50,3.40])
p('All 12 tests passed in the retained assessment. The decoder’s test descriptions use the word shape, but the helpers validate element count rather than every tensor axis. The authorization tests exercise policy helpers and do not execute a complete signed-session/database access path.')
minor('8.2 Test design principles')
p('Valid and invalid paths need separate evidence. A successful response demonstrates one execution path, while rejection cases establish specific constraints. Tests must assert the intended error code as well as status, because a generic server failure can otherwise hide a validation regression. Data and expectations should be fixed so later runs can distinguish implementation changes from input variation.')
p('Patient tests should use isolated synthetic organizations and assert that cross-tenant resources are unavailable. Database transaction tests should verify that patient, consent and audit records commit together or roll back together. Inference tests should include unusual channel layouts, missing model files, concurrent cold starts and malformed outputs. These are proposed additions, not recorded completed checks.')

head('8.3 Integration checks and requirement traceability',2,True)
table('Recorded local HTTP checks',['Input or condition','Expected observation','Result'],[['Missing image','400 MISSING_IMAGE','Pass'],['Unsupported mode','400 INVALID_MODE','Pass'],['Empty file','400 EMPTY_FILE','Pass'],['Non-image content','415 UNSUPPORTED_IMAGE_CONTENT','Pass'],['MIME mismatch','400 MIME_CONTENT_MISMATCH','Pass'],['Corrupt PNG','400 CORRUPT_IMAGE','Pass'],['One byte above 10 MiB','413 FILE_TOO_LARGE','Pass'],['Constant-color RGB image, binary','200 successful binary response','Pass'],['Same RGB image, severity','200 successful severity response','Pass']],[2.00,3.32,.68])
p('The generated input was a 512 × 512 RGB image containing no retinal structures. Both model calls returned structured results and model metadata. Their assigned labels are not clinical findings and are deliberately excluded from diagnostic-performance reporting. The record’s individual elapsed times are not a latency benchmark: no controlled repeated sample, percentile analysis or workload protocol was performed.')
table('Requirements mapped to available evidence',['Requirement','Evidence','Coverage limitation'],[['FR-01, FR-02','Signature unit tests and seven HTTP rejection checks','Pixel-limit and transport behavior not comprehensively tested.'],['FR-03, FR-04','Decoder tests and two synthetic inference checks','No training-parity or retinal-image evaluation.'],['FR-05, FR-06','Typed response structures, stored JSON, registry and UI review','No approval enforcement test; gate absent.'],['FR-07','Policy/input unit tests and source inspection','Complete database/session integration not verified.'],['FR-08 to FR-10','Schema and source review','Required workflow remains incomplete.']],[1.15,2.55,2.30])

head('9 Results and critical evaluation')
minor('9.1 Engineering results')
p('The project demonstrates two executable local model paths behind one research interface. The implementation distinguishes task-specific preprocessing, checks uploaded content and normalizes validated logits. Model metadata make artifact identity and known limitations inspectable. The patient DAL supplies a separate organization-scoped foundation with selected transactional audit writes.')
p('The recorded unit and HTTP checks passed, as did type checking and linting. These results support the specific behaviors tested. They do not establish model discrimination, calibration, patient safety, complete access-control coverage, database durability, scalability or production availability. A production build was not included in the retained manuscript verification results and is not reported as passed.')
table('Results interpretation',['Observation','Supported conclusion','Unsupported inference'],[['12 unit checks passed','Selected helpers satisfy their assertions.','All application security is verified.'],['Nine HTTP checks passed','Selected route cases behave as expected.','All malformed images and attacks are handled.'],['Both artifacts execute','Runtime integration works on the synthetic fixture.','DR diagnosis is accurate.'],['Model identity returned','Responses carry registry version metadata.','Artifact integrity is enforced at runtime.'],['Patient DAL exists','Record access is organized behind server checks.','Screening persistence and clinical care are complete.']],[1.65,2.13,2.22])
minor('9.2 Scientific limitations')
p('No supplied dataset manifest identifies training images, patient-level partitions, class balance, acquisition devices or expert grading. There is no verified export history, training configuration, independent test set, threshold-selection record or calibration study. Accuracy, sensitivity, specificity, macro-F1, AUROC and QWK therefore cannot be estimated from the available project evidence. Confusion matrices, learning curves and patient examples would be fabricated if added to this report.')
p('A non-retinal image can reach the model after technical validation. That outcome is evidence of a missing content/gradability gate, not evidence that the assigned label is meaningful. The weighted severity index and maximum score cannot determine referral urgency. No lesion-localization or explainability method is validated in the repository.')

head('9.3 Risk assessment and corrective priorities',2,True)
table('Engineering risk register',['Risk','Consequence','Proposed mitigation'],[['Unverified labels and preprocessing','Scores may be assigned incorrect meanings.','Obtain export/training contracts; add parity tests.'],['No retinal quality gate','Non-retinal or ungradable input receives labels.','Introduce validated content and gradability assessment.'],['Public inference route','Unauthenticated requests can consume resources.','Enforce identity, approval, consent and request limits.'],['Incomplete persistence','Outputs lack patient/exam and laterality context.','Associate immutable image/model records within an authorized workflow.'],['Color-space and tensor assumptions','Valid files may fail or be misprocessed.','Explicit color conversion and tensor-axis validation.'],['Limited load controls','Memory or CPU contention can degrade service.','Measure load; introduce queue, timeout and worker limits.'],['Unproven operations','Recovery or rollback may fail when needed.','Exercise backups, restoration, monitoring and version rollback.']],[1.40,2.0,2.60])
minor('9.4 Privacy and ethical scope')
p('The reported technical work used software fixtures and a generated constant-color image. It did not recruit participants, inspect identifiable patient records or evaluate patient outcomes. The screenshots show the application interface without clinical cases. Any future patient-based study must establish appropriate dataset access, ethics review, consent and retention arrangements before collection or use.')
p('The schema’s audit entity does not itself make an audit log append-only. Database permissions, retention rules and operational controls determine that property. Likewise, a consent record does not automatically authorize all inference; enforcement must occur in the relevant server path. Compliance with any regulatory framework is not claimed.')
p('Priority should first be given to model provenance and server-enforced boundaries. Additional interface features should follow those foundations, because a more complete-looking application can otherwise increase confidence without increasing the evidence behind its results.')

head('10 Installation and operation')
minor('10.1 Local setup sequence')
p('Use an isolated development environment and synthetic records. These commands describe the repository’s workflow; they are not a record of deployment performed while producing this report. Install a Node.js release supported by the pinned dependencies, npm and PostgreSQL 15 or newer. Keep the two ONNX artifacts in the models directory.')
code('npm ci\nnpm run db:validate\nnpm run db:generate')
p('Before database commands, configure DATABASE_URL in a local environment file using your own database and credentials. Use .env.example as the source of variable names. Do not commit secrets. The installation postinstall hook generates Prisma Client, so configuration needed by the project must be available during installation.')
code('npm run db:migrate -- --name initial_clinical_schema\nnpm run db:seed')
p('Apply development migrations only to the intended development database. Inspect existing migration state before choosing a new migration name. The explicit seed creates synthetic records and leaves both models blocked pending validation. Seeding is not a patient-data import process.')
p('For synthetic sign-in, provide a unique SESSION_SECRET of at least 32 characters and explicitly enable synthetic authentication in the development environment. The subject must match the seeded development identity. This provider is disabled automatically in production.')
code('ENABLE_SYNTHETIC_AUTH="true"\nSYNTHETIC_AUTH_SUBJECT="synthetic|clinician-reviewer"\nnpm run dev')
p('Open http://localhost:3000. Use /screening for research inference, /model-information to inspect model limitations, and /sign-in followed by /patients for synthetic patient workflows. A DATABASE_URL alone does not establish a complete clinical service.')
minor('10.2 Verification and build commands')
code('npm test\nnpm run typecheck\nnpm run lint\nnpm run build\nnpm run start')
p('The build regenerates Prisma Client before compilation. A deployment must include native-library compatibility and readable model artifacts. Production build/start commands are provided for reproducibility; the retained verification evidence does not certify a production release.')

head('10.3 Proposed deployment architecture',2,True)
p('Figure 8 is a future deployment target. It separates request handling, patient data and inference resource management. The current prototype does not provide the complete arrangement shown. Component selection, hosting region, capacity and cost require a separate operational design decision.')
image(HERE/'deployment.png','Proposed deployment target with explicit missing integrations',5.7)
p('The first deployment milestone should enforce model approval, user identity and patient/examination authorization at the server boundary. A validated quality gate should precede disease inference. Private object storage should hold original images and derived assets with content hashes and controlled access. Monitoring should distinguish request failures from model-validation status without logging patient data.')

head('10.4 Operations and troubleshooting',2,True)
table('Operational symptoms and investigation',['Symptom','Likely area','Engineering action'],[['Model fails on first request','Artifact path, native runtime or session initialization','Check local files and compatible runtime; preserve error logs without clinical payloads.'],['File rejected before inference','Signature, MIME, size or decode','Compare declared type with actual bytes and inspect the returned code.'],['Patient list is unavailable','Database, session or active membership','Verify local setup and server actor; do not bypass authorization.'],['No stored screening history','Missing workflow integration','Expected limitation; do not assume a database write occurred.'],['Review or follow-up route is not found','Disabled module flags','Expected current behavior; schema presence does not activate the feature.'],['Status remains blocked pending validation','Missing model evidence','Supply and review validation records; do not change a badge to claim readiness.']],[1.55,1.73,2.72])
minor('Release and maintenance procedure')
p('A future release should freeze source, dependency lockfile, model artifacts and preprocessing identities together. Re-run the relevant unit and integration suite against that package, verify model digests, and record environment details. Separate clinical-validation evidence from software-release evidence so a passing software build cannot substitute for evaluation of a changed model.')
p('Monitoring targets should include request count, failure-code distribution, cold and warm inference duration, memory use and queue depth where a queue exists. These are proposed measurements, not existing dashboards. Establish alert thresholds only after measuring a representative workload. Logs should avoid uploaded image bytes, patient identifiers and clinical payloads.')
p('Backups and rollback require exercises rather than configuration alone. Verify database restoration into an isolated environment, document object-storage retention and test a version rollback that preserves historical image/model associations. An inference artifact should never silently replace the identity behind an existing prediction.')

head('11 Future work and conclusion')
minor('11.1 Staged development plan')
table('Proposed development stages and exit evidence',['Stage','Scope','Exit evidence'],[['1 Model provenance','Recover dataset, labels, preprocessing, export and artifact records.','Reviewed, versioned model contract and parity tests.'],['2 Enforced server boundary','Identity, model approval, consent, request limits and resource checks.','Negative authorization and abuse tests pass.'],['3 Persistent research workflow','Private images, examination association, immutable predictions and audit.','Transaction, isolation and recovery tests pass.'],['4 Independent model evaluation','Patient-level partitions, reference grading, calibration and external validation.','Prespecified metrics with uncertainty and subgroup analysis.'],['5 Clinical workflow study','Usability, handling of ungradable inputs and clinician review.','Appropriate approvals and prospective workflow evidence.']],[1.15,2.45,2.40])
p('A diagnostic evaluation must define the target population, unit of analysis and reference standard before analysis. Multiple images from one patient should remain in the same data partition. Thresholds and calibration must be fitted without using the independent test set. Report excluded and ungradable images as well as classification metrics so operational failure is not concealed [7, 8].')
minor('11.2 Conclusion')
p('OpthaEdge demonstrates an engineering integration of two ONNX inference contracts with task-specific preprocessing, input/output checks, model metadata and an understandable research interface. Its patient-record foundation introduces server-side organization and permission boundaries, while the data schema prepares for a later persistent screening workflow.')
p('The strongest available result is reproducible software behavior for the recorded test cases. The main unresolved issue is the absence of provenance-complete and independently evaluated model evidence, accompanied by missing production identity, quality, storage and approval enforcement. The project is therefore best understood as a research software foundation. Progress should be judged by verified model contracts, enforced system boundaries and reproducible evaluation, not by the apparent confidence of an individual output.')

head('References')
src=Document(str(ROOT/'docs/research_paper/OpthaEdge_Research_Paper.docx'))
refs=[x.text for x in src.paragraphs if x.text.startswith('[') and x.text[1:2].isdigit()]
for ref in refs:
    e=p(ref);e.paragraph_format.space_after=Pt(10)
    for r in e.runs:r.font.size=Pt(10.5)
p('Repository evidence: OpthaEdge source commit 0099145dc14547c30484798042b029fbf58df1eb; README.md; prisma/schema.prisma; application, inference, authorization and test modules cited in this report. HTTP observations are retained in docs/research_paper/work/system_verification.json. Reference metadata are retained from the previously verified manuscript.')

head('Appendix A Reproducibility and artifact identity')
table('Recorded verification environment',['Item','Recorded value'],[['Source commit','0099145dc14547c30484798042b029fbf58df1eb'],['Assessment date','29 September 2026'],['Platform','Windows (win32)'],['Node runtime used for stored HTTP checks','v25.5.0'],['Processor','Intel Core i7-10850H at 2.70 GHz; 12 logical processors'],['HTTP fixture','One generated 512 × 512 constant-color RGB PNG'],['Data scope','No fundus photograph or patient information'],['Evidence location','docs/research_paper/work/system_verification.json']],[2.3,3.7])
minor('Artifact digests')
p('Binary artifact: retinal_binary_model.onnx')
code('dfc49e28bb3927ec47691320a12ae95d9af86b64bc16dd1240d56903d798f0d0')
p('Severity artifact: inceptionv3.onnx')
code('a75ff338af12bb31b66e05403b4b8ab989ea7c2213aa42447642d6b348028836')
p('These digests identify the reviewed model files and matched the registry in the stored assessment. They do not provide training provenance or prove that the running application checks file integrity. Recompute them when reproducing the work or distributing a new package, and record any mismatch before further testing.')
minor('Reproduction record to retain')
for t in ['Exact source commit, dependency lockfile, runtime and operating system.','Artifact hashes, preprocessing versions and the test fixture generation procedure.','Raw command output and raw HTTP responses, including failures.','A distinction between cold-start and warm-session runs if timing is studied.','A separate record for any database, security, load or clinical evaluation beyond this baseline.']:bullet(t)

head('Appendix B Source traceability')
table('Claims mapped to primary project files',['Report claim','Primary source'],[['Validation order and HTTP errors','app/api/predict/route.ts'],['File signatures and MIME alias handling','lib/imageFileSignature.ts'],['Task-specific preprocessing','lib/preprocessor.ts'],['CPU provider and shared initialization','lib/modelLoader.ts'],['Float32 outputs and elapsed-time field','lib/inference.ts'],['Threshold, argmax and score rounding','lib/outputDecoding.ts'],['Artifact versions and unverified assumptions','lib/modelRegistry.ts'],['API response types','types/prediction.ts'],['Eight-hour signed development session','lib/auth/session.ts'],['Active user/membership reconstruction','lib/auth/authorization.ts'],['Patient scoping and transaction','data/patients.ts'],['Patient form validation and server action','lib/patients/validation.ts; app/patients/actions.ts'],['Deferred module behavior','config/modules.ts; review-queue/follow-ups page files'],['Schema relationships and unique keys','prisma/schema.prisma'],['Recorded software checks','tests/*.test.ts; work/system_verification.json']],[2.70,3.30])
p('Historical design notes are useful context but may predate patient-record and authentication work. Where those notes conflict with current source, this report uses the current source and explicitly separates development functionality from production requirements. Earlier papers and all original supplied templates remain unchanged.')

# Materialize lists once captions exist; Word handles actual contents/page fields.
for item,entries,prefix in [(fig_list,figures,'fig'),(table_list,tables,'tab')]:
    item.clear()
    item.paragraph_format.tab_stops.add_tab_stop(Inches(6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    for i,(n,title) in enumerate(entries):
        if i:item.add_run().add_break()
        item.add_run(f'{n}. {title}\t');field(item,f' PAGEREF {prefix}{n} \\h ')
    item.paragraph_format.line_spacing=1.05;item.paragraph_format.space_after=Pt(5)
    for r in item.runs:r.font.size=Pt(10)
settings=d.settings.element
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
d.core_properties.title='OpthaEdge Engineering Project Report'
d.core_properties.subject='Diabetic retinopathy research software implementation and verification'
d.core_properties.author='Shamail Pawar; Faizan Mustafa Shaikh; Ayaan Shaikh'
d.core_properties.keywords='OpthaEdge, engineering report, diabetic retinopathy, ONNX, research software'
d.save(OUT)
report={'output':str(OUT),'figures':figures,'tables':tables,'references':len(refs),'paragraph_words':sum(len(x.text.split()) for x in d.paragraphs),'source_commit':'0099145dc14547c30484798042b029fbf58df1eb'}
(HERE/'build_inventory.json').write_text(json.dumps(report,indent=2))
print(json.dumps({k:v if k not in ['figures','tables'] else len(v) for k,v in report.items()},indent=2))
